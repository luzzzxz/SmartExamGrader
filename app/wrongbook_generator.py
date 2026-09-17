from __future__ import annotations

import json
import os
import re
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk


PROJECT_DIR = Path(__file__).resolve().parent
SETTINGS_PATH = PROJECT_DIR / "wrongbook_settings.json"
QUESTION_START_RE = re.compile(r"^\s*(?:\u7b2c\s*)?([0-9０-９]{1,3})\s*(?:[\u9898\u984c]|[.．\u3001,，])")
WORD_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_ATTRS = [
    f"{{{WORD_REL_NS}}}embed",
    f"{{{WORD_REL_NS}}}link",
    f"{{{WORD_REL_NS}}}id",
]


def load_settings():
    if not SETTINGS_PATH.exists():
        return {}
    try:
        return json.loads(SETTINGS_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_settings(settings):
    SETTINGS_PATH.write_text(json.dumps(settings, ensure_ascii=False, indent=2), encoding="utf-8")


def save_docx_safely(doc, out_path):
    out_path = Path(out_path)
    try:
        doc.save(str(out_path))
        return out_path
    except PermissionError:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback = out_path.with_name(f"{out_path.stem}_新生成_{stamp}{out_path.suffix}")
        doc.save(str(fallback))
        return fallback


def xml_tag_name(element):
    tag = getattr(element, "tag", "")
    return tag if isinstance(tag, str) else ""


def xml_local_name(element):
    tag = xml_tag_name(element)
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def max_plain_id(doc, local_names):
    max_id = 0
    for node in doc.element.iter():
        if xml_local_name(node) not in local_names:
            continue
        value = node.get("id")
        if value and str(value).isdigit():
            max_id = max(max_id, int(value))
    return max_id


def remove_node(node):
    parent = node.getparent()
    if parent is not None:
        parent.remove(node)


def unwrap_node(node):
    parent = node.getparent()
    if parent is None:
        return
    index = parent.index(node)
    for child in list(node):
        node.remove(child)
        parent.insert(index, child)
        index += 1
    parent.remove(node)


def clean_unstable_word_markup(element):
    # Comments and tracked-change wrappers can reference document-level parts
    # that are not copied into the merged document. Keep visible text, drop the
    # unstable markers so Word does not need to repair the generated file.
    remove_names = {
        "commentRangeStart",
        "commentRangeEnd",
        "commentReference",
        "del",
        "proofErr",
        "permStart",
        "permEnd",
    }
    unwrap_names = {"ins"}
    for node in list(element.iter()):
        name = xml_local_name(node)
        if name in remove_names:
            remove_node(node)
        elif name in unwrap_names:
            unwrap_node(node)


def renumber_drawing_ids(element, target_doc):
    next_docpr_id = max_plain_id(target_doc, {"docPr"}) + 1
    next_cnvpr_id = max_plain_id(target_doc, {"cNvPr"}) + 1
    for node in element.iter():
        name = xml_local_name(node)
        if name == "docPr" and node.get("id") is not None:
            node.set("id", str(next_docpr_id))
            next_docpr_id += 1
        elif name == "cNvPr" and node.get("id") is not None:
            node.set("id", str(next_cnvpr_id))
            next_cnvpr_id += 1


def image_work_dir(app):
    image_paths = getattr(app, "image_paths", None) or []
    parents = []
    for image_path in image_paths:
        try:
            path = Path(image_path)
        except Exception:
            continue
        if path.exists():
            parents.append(path.parent)
    if not parents:
        return None
    unique_parents = sorted({parent.resolve() for parent in parents})
    if len(unique_parents) == 1:
        return unique_parents[0]
    try:
        return Path(os.path.commonpath([str(parent) for parent in unique_parents]))
    except Exception:
        return unique_parents[0]


def default_work_dir(app):
    current_image_dir = image_work_dir(app)
    if current_image_dir:
        return current_image_dir
    for attr in ("selected_folder_path", "last_open_dir", "default_export_dir"):
        value = getattr(app, attr, "")
        if value:
            return Path(value)
    return PROJECT_DIR


def source_session_name(app):
    session_name = (getattr(app, "current_session", None) or {}).get("name", "")
    if session_name:
        return session_name
    work_dir = default_work_dir(app)
    if work_dir.parent and work_dir.parent != work_dir:
        return f"{work_dir.parent.name}\\{work_dir.name}"
    return work_dir.name


def format_wrong_item(item):
    qno = item.get("question_no", "")
    return f"第{qno}题"


def normalize_question_no(value):
    text = str(value or "").strip()
    table = str.maketrans("０１２３４５６７８９", "0123456789")
    text = text.translate(table)
    try:
        return int(text)
    except Exception:
        return text


def normalize_parent_question_no(value):
    text = str(value or "").strip()
    table = str.maketrans("０１２３４５６７８９", "0123456789")
    text = text.translate(table)
    match = re.match(r"\s*(\d{1,3})", text)
    if match:
        return int(match.group(1))
    return normalize_question_no(text)


def question_no_from_label(value):
    text = str(value or "").strip()
    table = str.maketrans("０１２３４５６７８９", "0123456789")
    text = text.translate(table)
    match = re.search(r"(?:选择题|第\s*)?(\d{1,3})\s*题", text)
    if match:
        return int(match.group(1))
    return ""


def collector_record_question_no(record):
    label = record.get("question_label", "") or record.get("题目标签", "")
    qno = question_no_from_label(label)
    if qno != "":
        return qno
    return normalize_parent_question_no(record.get("question_no", record.get("题号", "")))


def question_sort_key(value):
    qno = normalize_parent_question_no(value)
    return (0, qno) if isinstance(qno, int) else (1, str(qno))


def safe_filename(value):
    name = re.sub(r'[\\/:*?"<>|]+', "_", str(value or "").strip())
    return name.strip(" .") or "未命名"


def extract_question_no(text):
    match = QUESTION_START_RE.match(text or "")
    if not match:
        return None
    return normalize_question_no(match.group(1))


def iter_document_blocks(doc):
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = doc.element.body
    for child in body.iterchildren():
        tag = xml_tag_name(child)
        if tag.endswith("}p"):
            yield "paragraph", Paragraph(child, doc)
        elif tag.endswith("}tbl"):
            yield "table", Table(child, doc)


def block_text(block_type, block):
    if block_type == "paragraph":
        return block.text or ""
    rows = []
    for row in block.rows:
        cells = [" ".join(p.text for p in cell.paragraphs).strip() for cell in row.cells]
        rows.append(" ".join(cell for cell in cells if cell))
    return "\n".join(row for row in rows if row)


def append_block_element(target_doc, source_doc, source_element):
    copied = deepcopy(source_element)
    clean_unstable_word_markup(copied)
    renumber_drawing_ids(copied, target_doc)
    remap_relationships(copied, source_doc, target_doc)

    body = target_doc.element.body
    last_tag = xml_tag_name(body[-1]) if len(body) else ""
    if last_tag.endswith("}sectPr"):
        body.insert(len(body) - 1, copied)
    else:
        body.append(copied)


def remap_relationships(element, source_doc, target_doc):
    source_part = source_doc.part
    target_part = target_doc.part
    rid_map = {}

    for node in element.iter():
        for attr in REL_ATTRS:
            rid = node.get(attr)
            if not rid or rid not in source_part.rels:
                continue
            if rid not in rid_map:
                rel = source_part.rels[rid]
                if rel.is_external:
                    rid_map[rid] = target_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                elif rel.reltype.endswith("/image") and rid in source_part.related_parts:
                    image_part = source_part.related_parts[rid]
                    try:
                        rid_map[rid], _image = target_part.get_or_add_image(BytesIO(image_part.blob))
                    except Exception:
                        # Some Word images, such as EMF/WMF-like objects, are valid in Word
                        # but cannot be parsed by python-docx. Preserve the original part
                        # relationship instead of dropping the whole question.
                        rid_map[rid] = target_part.relate_to(image_part, rel.reltype)
                else:
                    rid_map[rid] = target_part.relate_to(source_part.related_parts[rid], rel.reltype)
            node.set(attr, rid_map[rid])


def collector_mark_records_from_entry(entry, app=None):
    records = entry.get("selected_records") or []
    if records:
        return records
    records = entry.get("标记题目") or []
    if records:
        return records
    if app is not None and hasattr(app, "build_collector_mark_records"):
        try:
            return app.build_collector_mark_records(entry) or []
        except Exception:
            return []
    return []


def wrongbook_items_for_entry(entry, app=None):
    wrong_by_qno = {}
    for wrong in entry.get("wrong_questions", []) or []:
        qno = normalize_parent_question_no(wrong.get("question_no", ""))
        if qno == "" or qno in wrong_by_qno:
            continue
        wrong_by_qno[qno] = {
            "question_no": qno,
            "question_type": wrong.get("question_type", ""),
            "source": "grading",
        }
    if wrong_by_qno:
        return sorted(wrong_by_qno.values(), key=lambda item: question_sort_key(item["question_no"]))

    for record in collector_mark_records_from_entry(entry, app):
        if record.get("是否标记") is False:
            continue
        qno = collector_record_question_no(record)
        if qno == "" or qno in wrong_by_qno:
            continue
        wrong_by_qno[qno] = {
            "question_no": qno,
            "question_type": record.get("question_type", record.get("题型", "")) or "采集",
            "source": "collector",
            "display_text": record.get("display_text", "") or record.get("标记内容", "") or record.get("question_label", record.get("题目标签", "")),
        }
    return sorted(wrong_by_qno.values(), key=lambda item: question_sort_key(item["question_no"]))


def student_field(entry, english_key, chinese_key):
    value = entry.get(english_key, "")
    if value in (None, ""):
        value = entry.get(chinese_key, "")
    return value


def iter_summary_entries(summary_data):
    if isinstance(summary_data, dict):
        for key in ("students", "summary_data", "学生结果"):
            value = summary_data.get(key)
            if isinstance(value, list):
                return value
        return []
    return summary_data or []


def collect_wrongbook_students(summary_data, app=None):
    students = []
    for entry in iter_summary_entries(summary_data):
        if not isinstance(entry, dict):
            continue
        wrong_items = wrongbook_items_for_entry(entry, app)
        if wrong_items:
            students.append({
                "score_id": student_field(entry, "score_id", "学号"),
                "student_name": student_field(entry, "student_name", "姓名"),
                "file": student_field(entry, "file", "文件"),
                "wrong_count": len(wrong_items),
                "wrong_text": "；".join(format_wrong_item(item) for item in wrong_items),
                "wrong_items": wrong_items,
            })
    return students


def collect_class_wrong_questions(summary_data, app=None):
    buckets = {}
    for entry_index, entry in enumerate(iter_summary_entries(summary_data)):
        if not isinstance(entry, dict):
            continue
        score_id = str(student_field(entry, "score_id", "学号") or "").strip()
        student_name = str(student_field(entry, "student_name", "姓名") or "").strip()
        file_name = str(student_field(entry, "file", "文件") or "").strip()
        student_key = score_id or student_name or file_name or f"row-{entry_index}"
        seen_in_student = set()

        for wrong in wrongbook_items_for_entry(entry, app):
            qno = normalize_parent_question_no(wrong.get("question_no", ""))
            if qno == "" or qno in seen_in_student:
                continue
            seen_in_student.add(qno)

            bucket = buckets.setdefault(qno, {
                "question_no": qno,
                "wrong_count": 0,
                "students": [],
                "_student_keys": set(),
            })
            if student_key in bucket["_student_keys"]:
                continue
            bucket["_student_keys"].add(student_key)
            bucket["wrong_count"] += 1
            bucket["students"].append({
                "score_id": score_id,
                "student_name": student_name,
                "file": file_name,
            })

    rows = []
    for bucket in buckets.values():
        clean_bucket = dict(bucket)
        clean_bucket.pop("_student_keys", None)
        rows.append(clean_bucket)
    return sorted(rows, key=lambda item: question_sort_key(item["question_no"]))


def split_docx_questions(paper_docx):
    from docx import Document

    paper_path = Path(paper_docx)
    if not paper_path.exists():
        raise FileNotFoundError(f"未找到Word试卷：{paper_path}")

    doc = Document(str(paper_path))
    questions = {}
    current_qno = None
    current_blocks = []
    current_texts = []

    def flush_current():
        if current_qno is None or not current_blocks:
            return
        questions[current_qno] = {
            "question_no": current_qno,
            "texts": list(current_texts),
            "blocks": list(current_blocks),
            "source_doc": doc,
        }

    for block_type, block in iter_document_blocks(doc):
        text = block_text(block_type, block).strip()
        if not text:
            if current_qno is not None and current_blocks:
                current_blocks.append(block._element)
                current_texts.append("")
            continue

        qno = extract_question_no(text)
        if qno is not None:
            flush_current()
            current_qno = qno
            current_blocks = [block._element]
            current_texts = [text]
        elif current_qno is not None:
            current_blocks.append(block._element)
            current_texts.append(text)

    flush_current()
    return questions


def build_question_preview(questions, max_items=30):
    if not questions:
        return "没有识别到题号。第一版支持形如 1.、1、和 第1题 的正文段落。"

    lines = [f"已识别 {len(questions)} 道题：", ""]
    for qno in sorted(questions.keys(), key=lambda value: value if isinstance(value, int) else 9999)[:max_items]:
        first_line = next((part for part in questions[qno].get("texts", []) if part.strip()), "")
        short = first_line[:80] + ("..." if len(first_line) > 80 else "")
        lines.append(f"第{qno}题：{short}")
    if len(questions) > max_items:
        lines.append(f"... 还有 {len(questions) - max_items} 道题未显示")
    return "\n".join(lines)


def add_question_to_doc(doc, qno, question, show_heading=True):
    if show_heading:
        doc.add_paragraph(f"第{qno}题", style="Heading 2")
    if not question:
        doc.add_paragraph(f"未在绑定的Word试卷中找到第{qno}题。")
        return
    source_doc = question.get("source_doc")
    if source_doc is not None and question.get("blocks"):
        for block in question.get("blocks", []):
            append_block_element(doc, source_doc, block)
        return

    for part in question.get("texts", []):
        doc.add_paragraph(part if part else "")


def generate_wrongbooks(app, paper_docx):
    from docx import Document

    students = collect_wrongbook_students(getattr(app, "summary_data", []), app)
    if not students:
        raise RuntimeError("当前没有可生成错题本的错题数据。")

    questions = split_docx_questions(paper_docx)
    export_dir = default_work_dir(app) / "错题本"
    export_dir.mkdir(parents=True, exist_ok=True)

    output_paths = []
    for student in students:
        doc = Document()
        name = student.get("student_name") or "未命名"
        score_id = student.get("score_id") or "未识别"
        doc.add_heading(f"{name} 错题本", level=1)
        doc.add_paragraph(f"测试/调查：{source_session_name(app)}")
        doc.add_paragraph(f"学号：{score_id}")
        doc.add_paragraph(f"错题：{student['wrong_text']}")

        for item in student["wrong_items"]:
            qno = item["question_no"]
            add_question_to_doc(doc, qno, questions.get(qno), show_heading=False)

        out_name = f"{safe_filename(score_id)}_{safe_filename(name)}_错题本.docx"
        out_path = export_dir / out_name
        output_paths.append(save_docx_safely(doc, out_path))

    return output_paths, questions, students


def generate_class_wrongbook(app, paper_docx, min_wrong_count=15):
    from docx import Document

    class_rows = collect_class_wrong_questions(getattr(app, "summary_data", []), app)
    selected_rows = [row for row in class_rows if row["wrong_count"] >= min_wrong_count]
    if not selected_rows:
        raise RuntimeError(f"没有达到“错误人数大于等于 {min_wrong_count}”的题目。")

    questions = split_docx_questions(paper_docx)
    export_dir = default_work_dir(app) / "错题本"
    export_dir.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("班级总体错题本", level=1)
    doc.add_paragraph(f"测试/调查：{source_session_name(app)}")

    for row in selected_rows:
        qno = row["question_no"]
        add_question_to_doc(doc, qno, questions.get(qno), show_heading=False)

    out_name = f"班级总体错题本_错题大于等于{min_wrong_count}人.docx"
    out_path = export_dir / out_name
    out_path = save_docx_safely(doc, out_path)
    return out_path, selected_rows, questions


def source_label_from_path(source_path, root_dir=None):
    path = Path(source_path)
    test_dir = path.parent.parent if path.parent.name == "错题本" else path.parent
    test_name = test_dir.name
    if test_name.endswith("错题本"):
        test_name = test_name[:-3] or test_dir.name
    class_name = test_dir.parent.name if test_dir.parent and test_dir.parent != test_dir else ""
    if root_dir:
        try:
            root = Path(root_dir).resolve()
            if test_dir.resolve() == root:
                class_name = test_dir.parent.name
        except Exception:
            pass
    return f"{class_name}\\{test_name}" if class_name else test_name


def should_skip_source_paragraph(text, skip_title=False):
    text = (text or "").strip()
    if not text:
        return False
    if text.startswith("测试/调查："):
        return True
    if skip_title and text in {"班级总体错题本", "合并班级总体错题本"}:
        return True
    return False


def copy_docx_content(target_doc, source_path, skip_metadata=False, skip_title=False):
    from docx import Document

    source_doc = Document(str(source_path))
    for block_type, block in iter_document_blocks(source_doc):
        if skip_metadata and block_type == "paragraph" and should_skip_source_paragraph(block.text, skip_title=skip_title):
            continue
        append_block_element(target_doc, source_doc, block._element)


def scan_generated_wrongbook_files(root_dir):
    root = Path(root_dir)
    files = list(root.rglob("*.docx")) if root.exists() else []
    student_files = []
    class_files = []
    for path in files:
        name = path.name
        if path.parent.name != "错题本":
            continue
        if name.startswith("合并") or "合并错题本" in name:
            continue
        if name.startswith("班级总体错题本"):
            class_files.append(path)
        elif name.endswith("_错题本.docx"):
            student_files.append(path)
    return sorted(student_files), sorted(class_files)


def student_merge_key(path):
    stem = path.stem
    if stem.endswith("_错题本"):
        stem = stem[:-4]
    return safe_filename(stem)


def merge_student_wrongbooks(app):
    from docx import Document

    work_dir = default_work_dir(app)
    student_files, _class_files = scan_generated_wrongbook_files(work_dir)
    if not student_files:
        raise RuntimeError("没有找到可合并的学生错题本。请先生成学生错题本，或选择包含错题本的目录。")

    groups = {}
    for path in student_files:
        groups.setdefault(student_merge_key(path), []).append(path)

    export_dir = work_dir / "错题本" / "合并学生错题本"
    export_dir.mkdir(parents=True, exist_ok=True)

    output_paths = []
    for key, paths in sorted(groups.items()):
        doc = Document()
        doc.add_heading(f"{key} 合并错题本", level=1)
        for index, path in enumerate(sorted(paths)):
            if index:
                doc.add_paragraph("")
            doc.add_paragraph(f"测试/调查：{source_label_from_path(path, work_dir)}")
            copy_docx_content(doc, path, skip_metadata=True)
        out_path = export_dir / f"{key}_合并错题本.docx"
        output_paths.append(save_docx_safely(doc, out_path))
    return output_paths, student_files


def merge_class_wrongbooks(app):
    from docx import Document

    work_dir = default_work_dir(app)
    _student_files, class_files = scan_generated_wrongbook_files(work_dir)
    if not class_files:
        raise RuntimeError("没有找到可合并的班级总体错题本。请先生成班级总体错题本，或选择包含错题本的目录。")

    export_dir = work_dir / "错题本"
    export_dir.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("合并班级总体错题本", level=1)
    for index, path in enumerate(class_files):
        if index:
            doc.add_paragraph("")
        doc.add_paragraph(f"测试/调查：{source_label_from_path(path, work_dir)}")
        copy_docx_content(doc, path, skip_metadata=True, skip_title=True)

    out_path = export_dir / "合并班级总体错题本.docx"
    out_path = save_docx_safely(doc, out_path)
    return out_path, class_files


def export_wrongbook_data(app, paper_docx):
    students = collect_wrongbook_students(getattr(app, "summary_data", []), app)
    export_dir = default_work_dir(app) / "错题本"
    export_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "paper_docx": paper_docx,
        "source_session": (getattr(app, "current_session", None) or {}).get("name", ""),
        "students": students,
    }
    out_path = export_dir / "错题本数据预览.json"
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return out_path, students


def open_wrongbook_window(app):
    settings = load_settings()
    students = collect_wrongbook_students(getattr(app, "summary_data", []), app)

    win = tk.Toplevel(app.root)
    win.title("错题本")
    win.geometry("880x600")
    win.minsize(820, 560)
    win.transient(app.root)

    paper_var = tk.StringVar(value=settings.get("paper_docx", ""))
    threshold_var = tk.StringVar(value=str(settings.get("class_wrong_threshold", 15)))
    summary_var = tk.StringVar(value=f"当前有错题学生 {len(students)} 人，错题 {sum(s['wrong_count'] for s in students)} 条")
    work_dir_var = tk.StringVar(value=f"当前目录：{default_work_dir(app)}")

    header = ttk.Frame(win)
    header.pack(fill="x", padx=16, pady=(14, 8))
    ttk.Label(header, text="Word错题本生成器", font=("Arial", 12, "bold")).pack(anchor="w")
    ttk.Label(
        header,
        text="按题号拆分Word题块，并尽量保留原卷的段落、图片、表格和公式结构。",
        foreground="#666",
    ).pack(anchor="w", pady=(4, 0))

    form = ttk.Frame(win)
    form.pack(fill="x", padx=16, pady=8)
    ttk.Label(form, text="Word试卷").grid(row=0, column=0, sticky="w", pady=6)
    paper_entry = ttk.Entry(form, textvariable=paper_var, width=72)
    paper_entry.grid(row=0, column=1, sticky="ew", padx=(8, 6), pady=6)
    form.columnconfigure(1, weight=1)

    def choose_paper():
        path = filedialog.askopenfilename(
            title="选择Word试卷",
            initialdir=str(default_work_dir(app)),
            filetypes=[("Word文档", "*.docx"), ("所有文件", "*.*")],
        )
        if not path:
            return
        paper_var.set(path)
        settings["paper_docx"] = path
        save_settings(settings)

    ttk.Button(form, text="选择", command=choose_paper).grid(row=0, column=2, sticky="e", pady=6)
    ttk.Label(form, text="班级错题阈值").grid(row=1, column=0, sticky="w", pady=6)
    ttk.Entry(form, textvariable=threshold_var, width=12).grid(row=1, column=1, sticky="w", padx=(8, 6), pady=6)
    ttk.Label(form, text="错误人数大于等于这个数时，抽入班级总体错题本", foreground="#666").grid(row=1, column=1, sticky="w", padx=(100, 6), pady=6)

    ttk.Label(win, textvariable=summary_var).pack(anchor="w", padx=16, pady=(4, 6))
    ttk.Label(win, textvariable=work_dir_var, foreground="#666").pack(anchor="w", padx=16, pady=(0, 6))

    preview = tk.Text(win, height=18, wrap="word")
    preview.pack(fill="both", expand=True, padx=16, pady=(0, 10))

    def set_preview(text):
        preview.config(state="normal")
        preview.delete("1.0", tk.END)
        preview.insert(tk.END, text)
        preview.config(state="disabled")

    def show_wrong_data_preview():
        lines = ["错题数据预览", ""]
        if not students:
            lines.append("当前还没有可用于错题本的错题数据。请先完成批改。")
        else:
            for student in students[:50]:
                lines.append(f"{student['score_id']} {student['student_name']}：{student['wrong_text']}")
            if len(students) > 50:
                lines.append(f"\n仅预览前50名学生，剩余 {len(students) - 50} 名会保留在数据中。")
        set_preview("\n".join(lines))

    show_wrong_data_preview()

    button_bar = ttk.Frame(win)
    button_bar.pack(fill="x", padx=16, pady=(0, 14))

    def save_binding():
        settings["paper_docx"] = paper_var.get().strip()
        settings["class_wrong_threshold"] = threshold_var.get().strip() or "15"
        save_settings(settings)
        if hasattr(app, "status_var"):
            app.status_var.set("已保存Word试卷绑定")
        win.destroy()

    def export_preview():
        path, exported_students = export_wrongbook_data(app, paper_var.get().strip())
        summary_var.set(f"已导出错题本数据预览：{path}（{len(exported_students)} 人）")
        if hasattr(app, "status_var"):
            app.status_var.set(f"已导出错题本数据预览：{path}")

    def preview_split():
        try:
            questions = split_docx_questions(paper_var.get().strip())
            set_preview(build_question_preview(questions))
            summary_var.set(f"Word试卷已识别 {len(questions)} 道题")
        except Exception as e:
            messagebox.showerror("拆题失败", str(e), parent=win)

    def build_wrongbooks():
        paper = paper_var.get().strip()
        if not paper:
            messagebox.showwarning("提示", "请先选择Word试卷。", parent=win)
            return
        try:
            paths, questions, exported_students = generate_wrongbooks(app, paper)
            summary_var.set(f"已生成 {len(paths)} 份错题本，识别试题 {len(questions)} 道")
            set_preview(
                "生成完成\n\n"
                + "\n".join(str(path) for path in paths[:80])
                + (f"\n... 还有 {len(paths) - 80} 个文件未显示" if len(paths) > 80 else "")
            )
            if hasattr(app, "status_var"):
                app.status_var.set(f"已生成 {len(paths)} 份错题本")
        except Exception as e:
            messagebox.showerror("生成失败", str(e), parent=win)

    def build_class_wrongbook():
        paper = paper_var.get().strip()
        if not paper:
            messagebox.showwarning("提示", "请先选择Word试卷。", parent=win)
            return
        try:
            threshold = int(threshold_var.get().strip() or "15")
            if threshold <= 0:
                raise ValueError
        except Exception:
            messagebox.showwarning("提示", "班级错题阈值请填写大于0的整数。", parent=win)
            return
        try:
            settings["paper_docx"] = paper
            settings["class_wrong_threshold"] = threshold
            save_settings(settings)
            path, selected_rows, questions = generate_class_wrongbook(app, paper, threshold)
            summary_var.set(f"已生成班级总体错题本：{len(selected_rows)} 道题，识别试题 {len(questions)} 道")
            lines = [
                "班级总体错题本生成完成",
                "",
                str(path),
                "",
                "抽取题目：",
            ]
            for row in selected_rows:
                lines.append(f"第{row['question_no']}题：{row['wrong_count']} 人做错")
            set_preview("\n".join(lines))
            if hasattr(app, "status_var"):
                app.status_var.set(f"已生成班级总体错题本：{path}")
        except Exception as e:
            messagebox.showerror("生成失败", str(e), parent=win)

    def build_merged_student_wrongbooks():
        try:
            paths, source_files = merge_student_wrongbooks(app)
            summary_var.set(f"已合并学生错题本：{len(paths)} 份，来源文件 {len(source_files)} 个")
            lines = [
                "学生错题本批量合并完成",
                "",
                f"输出目录：{default_work_dir(app) / '错题本' / '合并学生错题本'}",
                "",
                "生成文件：",
            ]
            for path in paths[:80]:
                lines.append(str(path))
            if len(paths) > 80:
                lines.append(f"... 还有 {len(paths) - 80} 个文件未显示")
            set_preview("\n".join(lines))
            if hasattr(app, "status_var"):
                app.status_var.set(f"已合并学生错题本：{len(paths)} 份")
        except Exception as e:
            messagebox.showerror("合并失败", str(e), parent=win)

    def build_merged_class_wrongbook():
        try:
            path, source_files = merge_class_wrongbooks(app)
            summary_var.set(f"已合并班级总体错题本：来源文件 {len(source_files)} 个")
            lines = [
                "班级总体错题本批量合并完成",
                "",
                str(path),
                "",
                "来源文件：",
            ]
            for source_file in source_files[:80]:
                lines.append(str(source_file))
            if len(source_files) > 80:
                lines.append(f"... 还有 {len(source_files) - 80} 个来源文件未显示")
            set_preview("\n".join(lines))
            if hasattr(app, "status_var"):
                app.status_var.set(f"已合并班级总体错题本：{path}")
        except Exception as e:
            messagebox.showerror("合并失败", str(e), parent=win)

    ttk.Button(button_bar, text="保存绑定", command=save_binding).pack(side="left", padx=(0, 8))
    ttk.Button(button_bar, text="拆题预览", command=preview_split).pack(side="left", padx=8)
    ttk.Button(button_bar, text="生成错题本", command=build_wrongbooks).pack(side="left", padx=8)
    ttk.Button(button_bar, text="生成班级总体错题本", command=build_class_wrongbook).pack(side="left", padx=8)
    ttk.Button(button_bar, text="合并学生错题本", command=build_merged_student_wrongbooks).pack(side="left", padx=8)
    ttk.Button(button_bar, text="合并班级错题本", command=build_merged_class_wrongbook).pack(side="left", padx=8)
    ttk.Button(button_bar, text="导出数据预览", command=export_preview).pack(side="left", padx=8)
    ttk.Button(button_bar, text="关闭", command=win.destroy).pack(side="right")

    if not students:
        messagebox.showinfo("提示", "当前还没有错题数据。可以先绑定Word试卷，批改后再回来生成错题本。", parent=win)
