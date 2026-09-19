# -*- coding: utf-8 -*-
from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _get_risk_level(risk_score: float, mastery: float):
    if risk_score >= 150 or mastery < 50:
        return '高危重点', '#ffebee', '#c62828'
    elif risk_score >= 60 or mastery < 80:
        return '需巩固', '#fff8e1', '#f57f17'
    else:
        return '良好', '#e8f5e9', '#2e7d32'


def export_class_report_html(
    target_path: Path,
    class_name: str,
    scope_desc: str,
    stat_mode: str,
    used_tests_count: int,
    rows: list[dict],
    weak_line: float = 80.0,
) -> Path:
    target_path.parent.mkdir(parents=True, exist_ok=True)
    generated_at = datetime.now().strftime('%Y-%m-%d %H:%M')

    table_rows_html = []
    high_risk_points = []

    for rank, row in enumerate(rows, 1):
        level_text, bg_color, text_color = _get_risk_level(row.get('risk_score', 0), row.get('mastery', 0))
        if level_text == '高危重点':
            high_risk_points.append(f"{row['code']} {row['name']}")

        table_rows_html.append(f"""
        <tr>
            <td style="text-align: center;">{rank}</td>
            <td style="text-align: center; font-weight: bold;">{row['code']}</td>
            <td>{row['name']}</td>
            <td style="text-align: center;">{row.get('weight', 1):g}</td>
            <td style="text-align: center;">{row.get('question_count', 0)}</td>
            <td style="text-align: center; color: #d32f2f; font-weight: bold;">{row.get('wrong_count', 0)}</td>
            <td style="text-align: center;">{row.get('score', 0):g} / {row.get('max_score', 0):g}</td>
            <td style="text-align: center; font-weight: bold;">{row.get('mastery', 0):.1f}%</td>
            <td style="text-align: center; font-weight: bold; color: {text_color};">{row.get('risk_score', 0):.0f}</td>
            <td style="text-align: center;"><span style="background: {bg_color}; color: {text_color}; padding: 2px 8px; border-radius: 4px; font-weight: bold; font-size: 12px;">{level_text}</span></td>
        </tr>
        """)

    summary_notice = ""
    if high_risk_points:
        summary_notice = f"""
        <div class="alert-box">
            <strong>⚠️ 班级高危薄弱重点（共 {len(high_risk_points)} 项，建议集中课堂讲评与专项强化）：</strong><br>
            <div style="margin-top: 6px; line-height: 1.6;">
                {'； '.join(high_risk_points[:12])}{' 等...' if len(high_risk_points) > 12 else '。'}
            </div>
        </div>
        """

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <title>{class_name} 知识点学情诊断分析报告</title>
    <style>
        @page {{
            size: A4 portrait;
            margin: 15mm 12mm 15mm 12mm;
        }}
        body {{
            font-family: "PingFang SC", "Microsoft YaHei", "SimHei", sans-serif;
            color: #333;
            margin: 0;
            padding: 20px;
            background: #fff;
        }}
        .header {{
            text-align: center;
            border-bottom: 2px solid #1976d2;
            padding-bottom: 12px;
            margin-bottom: 16px;
        }}
        .title {{
            font-size: 24px;
            font-weight: bold;
            color: #0d47a1;
            margin: 0 0 8px 0;
        }}
        .meta-info {{
            font-size: 13px;
            color: #555;
            display: flex;
            justify-content: space-between;
            flex-wrap: wrap;
            gap: 8px;
        }}
        .meta-item {{
            margin-right: 15px;
        }}
        .alert-box {{
            background-color: #fff3e0;
            border-left: 5px solid #ff9800;
            padding: 10px 14px;
            font-size: 13px;
            color: #e65100;
            margin-bottom: 16px;
            border-radius: 2px;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 13px;
        }}
        th, td {{
            border: 1px solid #ddd;
            padding: 7px 8px;
        }}
        th {{
            background-color: #f5f5f5;
            color: #333;
            font-weight: bold;
            text-align: center;
        }}
        tr:nth-child(even) {{
            background-color: #fafafa;
        }}
        .print-btn {{
            position: fixed;
            top: 20px;
            right: 20px;
            background: #1976d2;
            color: white;
            border: none;
            padding: 10px 18px;
            font-size: 14px;
            font-weight: bold;
            border-radius: 4px;
            cursor: pointer;
            box-shadow: 0 2px 8px rgba(0,0,0,0.2);
            z-index: 999;
        }}
        .print-btn:hover {{
            background: #0d47a1;
        }}
        @media print {{
            .print-btn {{
                display: none;
            }}
            body {{
                padding: 0;
            }}
        }}
    </style>
</head>
<body>
    <button class="print-btn" onclick="window.print()">🖨️ 打印本报告 (Ctrl+P)</button>
    <div class="header">
        <div class="title">{class_name} · 知识点学情综合诊断报告</div>
        <div class="meta-info">
            <span class="meta-item"><strong>班级：</strong>{class_name}</span>
            <span class="meta-item"><strong>统计范围：</strong>{scope_desc} (纳入 {used_tests_count} 次测试)</span>
            <span class="meta-item"><strong>统计模式：</strong>{stat_mode}</span>
            <span class="meta-item"><strong>生成时间：</strong>{generated_at}</span>
        </div>
    </div>

    {summary_notice}

    <table>
        <thead>
            <tr>
                <th style="width: 40px;">序号</th>
                <th style="width: 75px;">编号</th>
                <th>知识点内容</th>
                <th style="width: 45px;">权重</th>
                <th style="width: 65px;">作答人次</th>
                <th style="width: 65px;">失分人次</th>
                <th style="width: 80px;">得分/满分</th>
                <th style="width: 65px;">掌握率</th>
                <th style="width: 65px;">风险分</th>
                <th style="width: 75px;">综合研判</th>
            </tr>
        </thead>
        <tbody>
            {''.join(table_rows_html)}
        </tbody>
    </table>

    <div style="margin-top: 15px; font-size: 12px; color: #777; text-align: right;">
        * 风险分 = (100 - 掌握率) × 知识点权重。风险分越高，代表该高权重知识点在班级中失分越严重，为复习第一优先级。
    </div>
</body>
</html>
"""
    target_path.write_text(html, encoding='utf-8')
    return target_path


def export_class_report_excel(
    target_path: Path,
    class_name: str,
    scope_desc: str,
    stat_mode: str,
    used_tests_count: int,
    rows: list[dict],
) -> Path:
    target_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '班级知识点诊断'

    ws.merge_cells('A1:J1')
    title_cell = ws['A1']
    title_cell.value = f"{class_name} · 知识点学情综合诊断分析表"
    title_cell.font = Font(name='微软雅黑', size=16, bold=True, color='0D47A1')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 36

    ws.merge_cells('A2:J2')
    meta_cell = ws['A2']
    meta_cell.value = f"统计范围：{scope_desc} | 纳入测试：{used_tests_count} 场 | 统计模式：{stat_mode} | 导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    meta_cell.font = Font(name='微软雅黑', size=10, color='555555')
    meta_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 20

    headers = ['序号', '知识点编号', '知识点内容', '权重', '作答人次', '扣分人次', '总得分', '总满分', '掌握率', '风险分']
    ws.append(headers)
    ws.row_dimensions[3].height = 24

    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    header_fill = PatternFill(start_color='E3F2FD', end_color='E3F2FD', fill_type='solid')
    header_font = Font(name='微软雅黑', size=11, bold=True, color='1565C0')

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=3, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border

    for rank, row in enumerate(rows, 1):
        mastery = row.get('mastery', 0)
        risk = row.get('risk_score', 0)
        data_row = [
            rank,
            row.get('code', ''),
            row.get('name', ''),
            row.get('weight', 1),
            row.get('question_count', 0),
            row.get('wrong_count', 0),
            round(row.get('score', 0), 2),
            round(row.get('max_score', 0), 2),
            f"{mastery:.1f}%",
            round(risk, 1),
        ]
        ws.append(data_row)
        curr_row = 3 + rank
        ws.row_dimensions[curr_row].height = 20

        fill = None
        if risk >= 150 or mastery < 50:
            fill = PatternFill(start_color='FFEBEE', end_color='FFEBEE', fill_type='solid')
        elif risk >= 60 or mastery < 80:
            fill = PatternFill(start_color='FFF8E1', end_color='FFF8E1', fill_type='solid')

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.font = Font(name='微软雅黑', size=10)
            cell.border = thin_border
            if fill:
                cell.fill = fill
            if col_idx in (1, 2, 4, 5, 6, 7, 8, 9, 10):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    col_widths = [8, 14, 38, 8, 11, 11, 11, 11, 11, 11]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(target_path)
    return target_path


def export_student_worksheet_html_single(student_info: dict) -> str:
    student_name = student_info['student_name']
    score_id = student_info['score_id']
    class_name = student_info['class_name']
    scope_desc = student_info['scope_desc']
    stat_mode = student_info['stat_mode']
    rows = student_info['rows']
    generated_at = student_info.get('generated_at', datetime.now().strftime('%Y-%m-%d'))

    # 按系统加权风险分从高到低，固定取前 20 项核心考点，确保 A4 纸单面整齐装下
    display_rows = rows[:20]

    table_rows = []
    for rank, r in enumerate(display_rows, 1):
        level_text, bg_color, text_color = _get_risk_level(r.get('risk_score', 0), r.get('mastery', 0))
        wrong_details = [d for d in r.get('details', []) if d.get('score', 0) + 1e-9 < d.get('max_score', 0)]
        mistake_sources = []
        for wd in wrong_details[:2]:
            mistake_sources.append(f"{wd['session_name']} 第{wd['question_no']}题")
        mistake_sources_str = '；'.join(mistake_sources) if mistake_sources else '全对 / 无失分'

        table_rows.append(f"""
        <tr style="background-color: {'#fff9f9' if level_text == '高危重点' else '#ffffff'};">
            <td style="text-align: center;">{rank}</td>
            <td style="text-align: center; font-weight: bold;">{r['code']}</td>
            <td><strong>{r['name']}</strong></td>
            <td style="text-align: center;">{r.get('weight', 1):g}</td>
            <td style="text-align: center; font-weight: bold; color: {text_color};">{r.get('mastery', 0):.1f}%</td>
            <td style="text-align: center; font-weight: bold; color: {text_color};">{r.get('risk_score', 0):.0f}</td>
            <td style="font-size: 11px; color: {'#d32f2f' if wrong_details else '#2e7d32'};">{mistake_sources_str}</td>
            <td style="text-align: center;"><span style="background: {bg_color}; color: {text_color}; padding: 1px 6px; border-radius: 3px; font-weight: bold; font-size: 11px;">{level_text}</span></td>
            <td style="text-align: center; color: #9e9e9e;">[ &nbsp; ]</td>
        </tr>
        """)

    summary_bar = f"""
    <div style="display: flex; justify-content: space-between; align-items: center; margin: 4px 0 6px 0; font-size: 12px; background: #f8fafc; padding: 4px 10px; border-radius: 4px; border: 1px solid #e2e8f0;">
        <span>📊 本次共诊断 <strong>{len(rows)}</strong> 项考点，按<strong>综合风险分</strong>由高到低呈现重点 <strong>{len(display_rows)}</strong> 项：</span>
        <span>（重点攻关高危考点，订正错题）</span>
    </div>
    """

    if display_rows:
        urgent_table_html = f"""
        {summary_bar}
        <table>
            <thead>
                <tr>
                    <th style="width: 32px;">序号</th>
                    <th style="width: 58px;">编号</th>
                    <th>知识点名称</th>
                    <th style="width: 38px;">权重</th>
                    <th style="width: 58px;">掌握率</th>
                    <th style="width: 52px;">风险分</th>
                    <th>对应错题记录 / 表现</th>
                    <th style="width: 62px;">评级</th>
                    <th style="width: 52px;">已订正</th>
                </tr>
            </thead>
            <tbody>
                {''.join(table_rows)}
            </tbody>
        </table>
        """
    else:
        urgent_table_html = """
        <div style="background: #e8f5e9; color: #2e7d32; padding: 14px; border-radius: 4px; font-weight: bold; text-align: center; margin: 10px 0;">
            🎉 暂无作答数据记录。
        </div>
        """

    # 漏测/缺测知识点模块
    missing_rows = student_info.get('missing_rows', [])
    missing_table_html = ""
    if missing_rows:
        missing_tags = []
        for mr in missing_rows[:8]:
            missing_tags.append(f"<span style='display:inline-block; background:#fff3e0; color:#e65100; border:1px solid #ffe0b2; padding:2px 7px; border-radius:3px; margin:2px 4px 2px 0; font-size:11px;'><strong>{mr['code']}</strong> {mr['name']} (权重{mr.get('weight', 1):g})</span>")
        missing_tags_html = "".join(missing_tags)
        extra_txt = f" 等共 {len(missing_rows)} 项" if len(missing_rows) > 8 else ""

        missing_table_html = f"""
        <div style="margin-top: 8px; padding: 6px 10px; background: #fff8e1; border-left: 4px solid #ffb300; border-radius: 2px;">
            <div style="font-size: 12px; font-weight: bold; color: #b78103;">
                ⚠️ 缺测/漏测知识点（因缺考或考号填错，无近期测试数据，需重点自查或找老师过关）：
            </div>
            <div style="margin-top: 4px; line-height: 1.6;">
                {missing_tags_html}{extra_txt}
            </div>
        </div>
        """

    page_html = f"""
    <div class="student-page">
        <div class="student-header">
            <div class="student-title">学生个性化知识点复习指导单</div>
            <div class="student-meta">
                <span><strong>姓名：</strong><u>&nbsp;{student_name}&nbsp;</u></span>
                <span><strong>学号：</strong><u>&nbsp;{score_id or '无'}&nbsp;</u></span>
                <span><strong>班级：</strong>{class_name}</span>
                <span><strong>统计模式：</strong>{stat_mode}</span>
                <span><strong>诊断日期：</strong>{generated_at}</span>
            </div>
        </div>

        <div style="font-size: 12px; color: #555; margin-bottom: 6px; line-height: 1.4;">
            <strong>【复习指引】：</strong>本单汇总了你的近期作答诊断。按<strong>风险分高（失分严重/高权重）</strong>到低排列，请重点对照试卷订正标红错题。
        </div>

        {urgent_table_html}
        {missing_table_html}

        <div class="feedback-box">
            <div style="flex: 1; border-right: 1px dashed #ccc; padding-right: 15px;">
                <strong>💡 教师指导寄语：</strong><br>
                <div style="color: #666; font-size: 12px; margin-top: 4px;">
                    针对以上重点失分知识点，请在近期完成错题重做订正，有疑问随时向老师或同学请教！
                </div>
            </div>
            <div style="width: 260px; padding-left: 15px; font-size: 12px; display: flex; flex-direction: column; justify-content: space-around;">
                <div>学生自查签名：____________________</div>
                <div>家长复核签名：____________________</div>
            </div>
        </div>
    </div>
    """
    return page_html


def export_all_students_combined_html(
    target_path: Path,
    class_name: str,
    student_reports_data: list[dict],
) -> Path:
    target_path.parent.mkdir(parents=True, exist_ok=True)

    pages_html = []
    for s_info in student_reports_data:
        pages_html.append(export_student_worksheet_html_single(s_info))

    full_html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <title>{class_name} 全班学生复习指导单 (批量打印)</title>
    <style>
        @page {{
            size: A4 portrait;
            margin: 10mm 12mm 10mm 12mm;
        }}
        * {{
            box-sizing: border-box;
        }}
        body {{
            font-family: "PingFang SC", "Microsoft YaHei", "SimHei", sans-serif;
            color: #222;
            margin: 0;
            padding: 0;
            background: #f0f2f5;
        }}
        .print-toolbar {{
            position: fixed;
            top: 15px;
            right: 20px;
            background: rgba(25, 118, 210, 0.95);
            color: white;
            padding: 10px 18px;
            border-radius: 6px;
            box-shadow: 0 4px 12px rgba(0,0,0,0.25);
            z-index: 1000;
            display: flex;
            align-items: center;
            gap: 12px;
        }}
        .print-btn {{
            background: #fff;
            color: #1976d2;
            border: none;
            padding: 7px 16px;
            font-weight: bold;
            border-radius: 4px;
            cursor: pointer;
            font-size: 14px;
        }}
        .print-btn:hover {{
            background: #e3f2fd;
        }}
        .student-page {{
            background: #fff;
            width: 210mm;
            min-height: 285mm;
            padding: 10mm 12mm;
            margin: 15px auto;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
            position: relative;
            page-break-after: always;
            break-after: page;
        }}
        .student-header {{
            text-align: center;
            border-bottom: 2px solid #1565c0;
            padding-bottom: 4px;
            margin-bottom: 8px;
        }}
        .student-title {{
            font-size: 19px;
            font-weight: bold;
            color: #0d47a1;
            letter-spacing: 1px;
        }}
        .student-meta {{
            display: flex;
            justify-content: space-between;
            font-size: 11.5px;
            margin-top: 4px;
            color: #333;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 11px;
            margin-bottom: 6px;
        }}
        th, td {{
            border: 1px solid #ccc;
            padding: 3.5px 5px;
        }}
        th {{
            background: #f1f5f9;
            color: #1e293b;
            font-weight: bold;
            text-align: center;
        }}
        .good-tag {{
            display: inline-block;
            background: #f1f8e9;
            color: #33691e;
            border: 1px solid #c5e1a5;
            padding: 1px 6px;
            border-radius: 3px;
            margin: 2px 3px 2px 0;
            font-size: 11px;
        }}
        .feedback-box {{
            margin-top: 8px;
            border: 1px solid #90caf9;
            background: #f8fbff;
            border-radius: 4px;
            padding: 6px 10px;
            display: flex;
        }}
        @media print {{
            .print-toolbar {{
                display: none !important;
            }}
            body {{
                background: none;
            }}
            .student-page {{
                margin: 0;
                padding: 0;
                box-shadow: none;
                width: 100%;
                min-height: auto;
                page-break-after: always;
                break-after: page;
            }}
        }}
    </style>
</head>
<body>
    <div class="print-toolbar">
        <span>全班共 <strong>{len(student_reports_data)}</strong> 名学生</span>
        <button class="print-btn" onclick="window.print()">🖨️ 打印全班复习单 (自动分页)</button>
    </div>

    {''.join(pages_html)}
</body>
</html>
"""
    target_path.write_text(full_html, encoding='utf-8')
    return target_path


def export_all_students_excel(
    target_path: Path,
    class_name: str,
    scope_desc: str,
    stat_mode: str,
    student_reports_data: list[dict],
) -> Path:
    target_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '全班学生薄弱知识点一览'

    ws.merge_cells('A1:H1')
    title_cell = ws['A1']
    title_cell.value = f"{class_name} · 全班学生薄弱知识点分析汇总表"
    title_cell.font = Font(name='微软雅黑', size=15, bold=True, color='0D47A1')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    ws.merge_cells('A2:H2')
    meta_cell = ws['A2']
    meta_cell.value = f"统计范围：{scope_desc} | 模式：{stat_mode} | 学生人数：{len(student_reports_data)} | 导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    meta_cell.font = Font(name='微软雅黑', size=10, color='666666')
    meta_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 20

    headers = ['学号', '姓名', '知识点编号', '知识点内容', '权重', '掌握率', '风险分', '评级']
    ws.append(headers)
    ws.row_dimensions[3].height = 22

    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    header_fill = PatternFill(start_color='E3F2FD', end_color='E3F2FD', fill_type='solid')
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=3, column=col_idx)
        c.fill = header_fill
        c.font = Font(name='微软雅黑', size=10, bold=True, color='1565C0')
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border

    curr_row = 3
    for s_info in student_reports_data:
        s_name = s_info['student_name']
        s_id = s_info['score_id']
        for r in s_info['rows']:
            if r.get('mastery', 100) >= 100 and r.get('risk_score', 0) <= 0:
                continue
            curr_row += 1
            mastery = r.get('mastery', 0)
            risk = r.get('risk_score', 0)
            level_text, _, _ = _get_risk_level(risk, mastery)

            row_data = [
                s_id,
                s_name,
                r.get('code', ''),
                r.get('name', ''),
                r.get('weight', 1),
                f"{mastery:.1f}%",
                round(risk, 1),
                level_text
            ]
            ws.append(row_data)
            ws.row_dimensions[curr_row].height = 19

            fill = None
            if level_text == '高危重点':
                fill = PatternFill(start_color='FFEBEE', end_color='FFEBEE', fill_type='solid')
            elif level_text == '需巩固':
                fill = PatternFill(start_color='FFF8E1', end_color='FFF8E1', fill_type='solid')

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = Font(name='微软雅黑', size=9)
                cell.border = thin_border
                if fill:
                    cell.fill = fill
                if col_idx in (1, 2, 3, 5, 6, 7, 8):
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')

        # 记录漏测知识点
        for mr in s_info.get('missing_rows', []):
            curr_row += 1
            row_data = [
                s_id,
                s_name,
                mr.get('code', ''),
                mr.get('name', ''),
                mr.get('weight', 1),
                '未测',
                '-',
                '⚠️ 漏测'
            ]
            ws.append(row_data)
            ws.row_dimensions[curr_row].height = 19
            fill = PatternFill(start_color='FFF3E0', end_color='FFF3E0', fill_type='solid')

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = Font(name='微软雅黑', size=9)
                cell.border = thin_border
                cell.fill = fill
                if col_idx in (1, 2, 3, 5, 6, 7, 8):
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')

    col_widths = [12, 14, 14, 35, 8, 12, 12, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(target_path)
    return target_path
